"""Category screener: runs the metrics engine over every fund in a category and ranks peers.

Usage: python3 screener.py <universe_nav.csv> <benchmark_tri.csv> <benchmark name> <out.xlsx> [as_of]
"""
import sys
import warnings

import numpy as np
import pandas as pd

from metrics import (calendar_returns, max_drawdown, risk_metrics, rolling_returns, sip_returns,
                     trailing_returns, value_on, load_tri_csv, period_return)

warnings.filterwarnings("ignore")

# metrics where LOWER is better (rank ascending)
LOWER_BETTER = {"Std dev 3Y", "Beta 3Y", "Downside capture 3Y", "TER"}


def load_universe(path):
    df = pd.read_csv(path)
    df["date"] = pd.to_datetime(df["date"], format="%d-%m-%Y")
    meta = df.groupby("code").agg(amc=("amc", "first"), name=("name", "first")).reset_index()
    navs = {c: g.set_index("date")["nav"].astype(float).sort_index()[lambda s: ~s.index.duplicated()]
            for c, g in df.groupby("code")}
    return meta, navs


def short_name(n):
    return (n.split(" - ")[0].replace("Fund", "").replace("FUND", "").strip().title()
            .replace("Mf ", "MF ").replace("Sbi", "SBI").replace("Dsp", "DSP").replace("Hdfc", "HDFC")
            .replace("Iti", "ITI").replace("Lic", "LIC").replace("Pgim", "PGIM").replace("Hsbc", "HSBC")
            .replace("Icici", "ICICI").replace("Uti", "UTI").replace("Jm ", "JM ").replace("Trustmf", "TrustMF"))


def screen(meta, navs, bm, bm_name, as_of=None):
    as_of = pd.Timestamp(as_of) if as_of else max(s.index[-1] for s in navs.values())
    rows, cal_rows = [], []
    for _, m in meta.iterrows():
        s = navs[m.code].loc[:as_of]
        if len(s) < 20:
            continue
        d_last, nav_last = value_on(s, as_of)
        age = (d_last - s.index[0]).days / 365.25
        tr = trailing_returns(s, d_last)
        rk = risk_metrics(s, bm, d_last, 3)
        r3 = rolling_returns(s, bm, 3, d_last)
        if r3.get("observations", 0) < 250:  # need ~1 year of start dates to be meaningful
            r3 = {}
        sip3 = sip_returns(s, d_last, 3)
        dd = max_drawdown(s)
        has = "std_dev" in rk
        rows.append({
            "Code": m.code, "Fund": short_name(m["name"]), "AMC": m.amc,
            "First NAV": s.index[0].date(), "Age (yrs)": round(age, 1),
            "NAV": nav_last, "NAV date": d_last.date(),
            "1M": tr["1M"], "3M": tr["3M"], "6M": tr["6M"], "1Y": tr["1Y"], "3Y": tr["3Y"],
            "5Y": tr["5Y"], "10Y": tr["10Y"], "Since inception": tr["Since inception"],
            "SIP XIRR 3Y": sip3["xirr"] if sip3 else None,
            "Std dev 3Y": rk.get("std_dev"), "Sharpe 3Y": rk.get("sharpe"), "Sortino 3Y": rk.get("sortino"),
            "Beta 3Y": rk.get("beta"), "Alpha 3Y": rk.get("alpha"), "Info ratio 3Y": rk.get("information_ratio"),
            "Upside capture 3Y": rk.get("upside_capture"), "Downside capture 3Y": rk.get("downside_capture"),
            "Max DD 3Y": rk.get("max_drawdown") if has else None, "Max DD since start": dd["max_drawdown"],
            "Rolling 3Y avg": r3.get("average"), "Rolling 3Y % beat BM": r3.get("pct_beat_benchmark"),
            "Rolling 3Y % negative": r3.get("pct_negative"),
        })
        c = calendar_returns(s, d_last)
        cal_rows.append({"Fund": short_name(m["name"]), **{k: v for k, v in c.items() if "partial" not in k}})
    df = pd.DataFrame(rows)
    cal = pd.DataFrame(cal_rows).set_index("Fund")
    cal = cal[[c for c in cal.columns if c >= "2016"]]
    # benchmark row
    b_tr = trailing_returns(bm.loc[:as_of], as_of)
    bm_row = {"Fund": f"BENCHMARK: {bm_name}", **{k: b_tr[k] for k in ["1M", "3M", "6M", "1Y", "3Y", "5Y", "10Y"]}}
    b_cal = {}
    for col in cal.columns:
        y = int(col[:4])
        b_cal[col] = period_return(bm, pd.Timestamp(y - 1, 12, 31), min(pd.Timestamp(y, 12, 31), as_of), False)
    cal.loc[f"BENCHMARK: {bm_name}"] = pd.Series(b_cal)
    # peer ranks (1 = best) and quartiles, only among funds that have the metric
    rank_cols = ["1Y", "3Y", "5Y", "10Y", "Sharpe 3Y", "Sortino 3Y", "Alpha 3Y", "Info ratio 3Y",
                 "Std dev 3Y", "Downside capture 3Y", "Max DD 3Y", "Rolling 3Y % beat BM"]
    ranks = pd.DataFrame({"Fund": df.Fund})
    for c in rank_cols:
        asc = c in LOWER_BETTER
        if c == "Max DD 3Y":
            asc = False  # less negative is better
        r = df[c].rank(ascending=asc, method="min")
        n = df[c].notna().sum()
        ranks[c] = r.map(lambda x: f"{int(x)}/{n}" if pd.notna(x) else "")
        df[f"Q {c}"] = np.ceil(r / n * 4).where(df[c].notna())
    q3 = [c for c in df.columns if c.startswith("Q ") and c[2:] in
          ["3Y", "Sharpe 3Y", "Alpha 3Y", "Downside capture 3Y", "Rolling 3Y % beat BM"]]
    df["Avg quartile (3Y core)"] = df[q3].mean(axis=1).round(2)
    df = pd.concat([df, pd.DataFrame([bm_row])], ignore_index=True)
    return df, cal, ranks, as_of


def write_excel(df, cal, ranks, out, bm_name, as_of):
    pct_cols = ["1M", "3M", "6M", "1Y", "3Y", "5Y", "10Y", "Since inception", "SIP XIRR 3Y", "Std dev 3Y",
                "Alpha 3Y", "Upside capture 3Y", "Downside capture 3Y", "Max DD 3Y", "Max DD since start",
                "Rolling 3Y avg", "Rolling 3Y % beat BM", "Rolling 3Y % negative"]
    main = df[[c for c in df.columns if not c.startswith("Q ")]].sort_values(
        "Avg quartile (3Y core)", na_position="last")
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        main.to_excel(xw, sheet_name="Screener", index=False, startrow=2)
        cal.reset_index().to_excel(xw, sheet_name="Calendar years", index=False, startrow=2)
        ranks.to_excel(xw, sheet_name="Peer ranks", index=False, startrow=2)
        notes = pd.DataFrame({"Note": [
            f"Data as of {as_of:%d-%b-%Y}. NAV: AMFI / mfapi.in (Direct Growth). Benchmark: {bm_name} (niftyindices.com).",
            "Trailing: <1Y absolute, >=1Y CAGR. Calendar: 31-Dec to 31-Dec.",
            "Risk ratios: 36 month-end returns to the last month-end; annualised; risk-free 6.0%.",
            "Capture ratios: geometric, Morningstar method. Downside capture < 100% = falls less than index.",
            "Rolling 3Y: daily start dates, CAGR; % of periods the fund beat the benchmark.",
            "Avg quartile (3Y core) = mean of quartiles on 3Y return, Sharpe, Alpha, Downside capture, Rolling % beat. 1 = best.",
            "Funds under 3 years old have no 3Y ratios and are not ranked on them. Rolling 3Y needs 4+ years of history.",
            "HSBC Small Cap: history only from Nov-2022 under current scheme code (pre-merger L&T history not linked yet).",
        ]})
        notes.to_excel(xw, sheet_name="Method", index=False)
        wb = xw.book
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        for ws, frame, title in [(wb["Screener"], main, "Small Cap Screener"),
                                 (wb["Calendar years"], cal.reset_index(), "Calendar-year returns"),
                                 (wb["Peer ranks"], ranks, "Peer ranks (1 = best)")]:
            ws["A1"] = f"{title} — as of {as_of:%d %b %Y} — benchmark {bm_name}"
            ws["A1"].font = Font(bold=True, size=13)
            for i, col in enumerate(frame.columns, 1):
                h = ws.cell(row=3, column=i)
                h.font = Font(bold=True, color="FFFFFF")
                h.fill = PatternFill("solid", fgColor="1F3864")
                h.alignment = Alignment(wrap_text=True, vertical="center")
                ws.column_dimensions[get_column_letter(i)].width = 30 if i in (2, 3) or col == "Fund" else 11
                is_pct = col in pct_cols or ws.title == "Calendar years" and col != "Fund"
                for r in range(4, 4 + len(frame)):
                    c = ws.cell(row=r, column=i)
                    if isinstance(c.value, (int, float)):
                        c.number_format = "0.0%" if is_pct else "0.00"
                        if col in ("NAV",):
                            c.number_format = "#,##0.00"
            ws.freeze_panes = "C4" if ws.title == "Screener" else "B4"
            ws.row_dimensions[3].height = 32
            for r in range(4, 4 + len(frame)):
                if str(ws.cell(row=r, column=2 if ws.title == "Screener" else 1).value).startswith("BENCHMARK"):
                    for cc in ws[r]:
                        cc.font = Font(bold=True)
                        cc.fill = PatternFill("solid", fgColor="FFF2CC")
        wb["Method"].column_dimensions["A"].width = 120


if __name__ == "__main__":
    uni, bm_path, bm_name, out = sys.argv[1:5]
    as_of = sys.argv[5] if len(sys.argv) > 5 else None
    meta, navs = load_universe(uni)
    bm = load_tri_csv(bm_path)
    df, cal, ranks, as_of = screen(meta, navs, bm, bm_name, as_of)
    write_excel(df, cal, ranks, out, bm_name, as_of)
    pd.set_option("display.width", 250, "display.max_columns", 30)
    show = ["Fund", "Age (yrs)", "1Y", "3Y", "5Y", "10Y", "Std dev 3Y", "Sharpe 3Y", "Beta 3Y", "Alpha 3Y",
            "Upside capture 3Y", "Downside capture 3Y", "Max DD 3Y", "Rolling 3Y % beat BM", "Avg quartile (3Y core)"]
    t = df[show].sort_values("Avg quartile (3Y core)", na_position="last").copy()
    for c in show[2:-1]:
        if c not in ("Sharpe 3Y", "Beta 3Y"):
            t[c] = (t[c] * 100).round(1)
        else:
            t[c] = t[c].round(2)
    print(t.to_string(index=False))
    print(cal.mul(100).round(1).to_string())
